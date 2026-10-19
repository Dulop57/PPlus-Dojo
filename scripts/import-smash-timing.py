"""Read the owner's checked smash timings; never modify the supplied workbook."""
import argparse
import hashlib
import json
import pathlib
import re

import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument("workbook", type=pathlib.Path)
args = parser.parse_args()
root = pathlib.Path(__file__).resolve().parent.parent
manifest = json.loads((root / "public/data/manifest.json").read_text())
fighters = {re.sub(r"[^a-z0-9]", "", f["name"].lower()): f for f in manifest["fighters"]}
fighters["mrgameandwatch"] = fighters["gamewatch"]
workbook = openpyxl.load_workbook(args.workbook, data_only=True)
sheet = workbook["Sheet1"]
result = {"source": args.workbook.name, "sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
          "sheet": sheet.title, "fighters": {}}


def windows(value):
    if value is None:
        return None
    ranges = []
    for part in str(value).replace("–", "-").replace("—", "-").split(","):
        pair = [int(n.strip()) for n in part.strip().split("-")]
        if len(pair) == 1:
            pair *= 2
        assert len(pair) == 2 and 0 < pair[0] <= pair[1]
        assert not ranges or ranges[-1][1] < pair[0], "Overlapping active windows"
        ranges.append(pair)
    return ranges


for row in range(3, sheet.max_row + 1):
    name = sheet.cell(row, 1).value
    if not name:
        continue
    fighter = fighters[re.sub(r"[^a-z0-9]", "", name.lower())]
    index = json.loads((root / "public/data" / fighter["path"]).read_text())
    entries = {}
    for column, family in [(2, "S4"), (4, "Hi4"), (6, "Lw4")]:
        active = windows(sheet.cell(row, column).value)
        fa = sheet.cell(row, column + 1).value
        if fa is None:
            continue
        actionables = [int(v.strip()) for v in str(fa).split(",")]
        double = (fighter["id"] == "ike" and family == "Lw4") or (fighter["id"] in ["link", "pit"] and family == "S4")
        assert len(actionables) == (2 if double else 1)
        for entry in index["moves"]:
            internal = entry["internalName"]
            normal = re.fullmatch(r"Attack" + family + r"(?:S(?:Hi|Lw)?|Hi|Lw)?(?:_\d+)?", internal)
            first = double and internal == ("AttackLw41" if family == "Lw4" else "AttackS4S")
            second = double and internal == ("AttackLw42" if family == "Lw4" else "AttackS4S2")
            if not entry["totalFrames"] or not (normal or first or second):
                continue
            # Underscore variants are alternate copies of the first attack, not the follow-up.
            second_phase = bool(second)
            desired = active if not double or second_phase else active[:1]
            prefix_name = ("AttackLw41" if family == "Lw4" else "AttackS4S") if second_phase else None
            dependencies = ["Attack" + family + "Start"] + ([prefix_name] if prefix_name else [])
            entries[internal] = {"sha256": entry["sha256"], "family": family, "active": desired,
                                 "actionable": actionables[1 if second_phase else 0],
                                 "cells": f"{openpyxl.utils.get_column_letter(column)}{row}:{openpyxl.utils.get_column_letter(column+1)}{row}",
                                 "dependencies": {e["internalName"]: e["sha256"] for e in index["moves"] if e["internalName"] in dependencies}}
            if prefix_name:
                entries[internal]["prefix"] = prefix_name
    result["fighters"][fighter["id"]] = entries
assert len(result["fighters"]) == 42
destination = root / "src/data/smash-timing.json"
destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Imported {len(result['fighters'])} fighters, {sum(map(len, result['fighters'].values()))} smash subactions -> {destination}")
