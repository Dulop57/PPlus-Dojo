"""Read supplied XLSX files without modifying them; export static reference data.

Run with Python + openpyxl. No game files are read by this importer.
"""
import argparse, datetime, hashlib, json
from pathlib import Path
import openpyxl

def value(v):
    return v.isoformat()[:10] if isinstance(v, (datetime.datetime, datetime.date)) else v

def workbook(path):
    original = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    missing = []
    for sheet in original:
        cells = []
        for row in sheet:
            for cell in row:
                if cell.value is None and not cell.comment:
                    continue
                result = {'r':cell.row, 'c':cell.column, 'v':value(cached[sheet.title][cell.coordinate].value)}
                if cell.data_type == 'f':
                    result['formula'] = cell.value
                    if result['v'] is None:
                        missing.append(f'{sheet.title}!{cell.coordinate}')
                if cell.font.bold: result['bold'] = True
                if cell.comment: result['note'] = cell.comment.text
                if cell.hyperlink and cell.hyperlink.target: result['link'] = cell.hyperlink.target
                if cell.number_format != 'General': result['format'] = cell.number_format
                cells.append(result)
        maxrow = max((c['r'] for c in cells), default=1)
        maxcol = max((c['c'] for c in cells), default=1)
        merges = [[m.min_row,m.min_col,m.max_row,m.max_col] for m in sheet.merged_cells.ranges if m.min_row <= maxrow and m.min_col <= maxcol]
        sheets.append({'name':sheet.title,'rows':maxrow,'columns':maxcol,'cells':cells,'merges':merges})
    return {'schemaVersion':1,'source':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':sheets,'missingCachedValues':missing}

def stages(path):
    raw = workbook(path)
    book = openpyxl.load_workbook(path, data_only=True)
    original = openpyxl.load_workbook(path, data_only=False)
    main, coords = book['Page 1'], book['Raw data']
    indexed = {str(coords.cell(r,1).value):r for r in range(2,coords.max_row+1) if coords.cell(r,1).value}
    result = []
    for row in range(2,main.max_row+1):
        name = main.cell(row,1).value
        if not name or not isinstance(main.cell(row,2).value,(float,int)): continue
        base = name.split('\n')[0]
        rr = indexed.get(base)
        assert rr, f'Missing coordinates for {base}'
        def c(col): return coords.cell(rr,col).value
        notes = [{'cell':cell.coordinate,'text':cell.comment.text} for cell in original['Page 1'][row] if cell.comment]
        result.append({'id':'-'.join(base.lower().replace("'",'').split()),'name':base,'variant':' '.join(name.split('\n')[1:]) or 'Standard',
            'summary':{'width':main.cell(row,2).value,'top':main.cell(row,3).value,'side':main.cell(row,4).value,'bottom':main.cell(row,5).value,'low':main.cell(row,6).value,'high':main.cell(row,7).value,'platformWidth':main.cell(row,8).value},
            'geometry':{'left':c(11),'right':c(12),'ledgeY':c(13),'blastLeft':c(16),'blastTop':c(17),'blastRight':c(20),'blastBottom':c(21),'platformLow':c(22),'platformHigh':c(23),'platformLeft':c(24),'platformRight':c(25)},
            'notes':main.cell(row,9).value,'cellNotes':notes,'sourceRow':row,'rawRow':rr})
    raw['stages'] = result
    return raw

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--attributes',type=Path,required=True)
    parser.add_argument('--stages',type=Path,required=True)
    parser.add_argument('--out',type=Path,default=Path('src/data/reference'))
    args=parser.parse_args(); args.out.mkdir(parents=True,exist_ok=True)
    for name,data in [('attributes',workbook(args.attributes)),('stages',stages(args.stages))]:
        target=args.out / f'{name}.json'
        target.write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
        print(f'{target}: {len(data["sheets"])} sheets; {len(data["missingCachedValues"])} unavailable formula caches; {target.stat().st_size} bytes')
    print('Original XLSX files were not modified.')

if __name__=='__main__':main()
